import asyncio
import os
import sys
import csv
from datetime import datetime
import openpyxl

# Add project root to sys.path
sys.path.insert(0, r"C:\Users\ASUS\.gemini\antigravity-ide\scratch\leadgen")

from module_1_scraping_engine.run_scraper import ARCHETYPE_MAP
from module_1_scraping_engine.pipelines.validator import JobValidator

async def run_all_archetype_e():
    scrapers = ARCHETYPE_MAP.get("e", [])
    print(f"Total Archetype E External Stealth Portals: {len(scrapers)}")
    
    all_records = []
    summary_stats = []
    
    fieldnames = [
        "Platform / Portal",
        "Source ID",
        "Job Title",
        "Company",
        "Location",
        "Is Remote",
        "Employment Type",
        "Salary (Scale / Raw)",
        "Salary Min",
        "Salary Max",
        "Salary Currency",
        "Apply URL",
        "Scraped Timestamp"
    ]
    
    for scraper_cls in scrapers:
        scraper_name = scraper_cls.__name__
        scraper = None
        try:
            scraper = scraper_cls()
            print(f"\n=======================================================")
            print(f"Running [{scraper.source_id}] {scraper.source_name}...")
            print(f"=======================================================")
            
            jobs = await scraper.scrape(limit=10)
            await scraper.close()
            
            # Validate through standard schema validator
            valid_jobs, rejected = JobValidator.validate_batch([j.model_dump() for j in jobs])
            
            count = len(valid_jobs)
            status = "SUCCESS" if count > 0 else "NO_DATA"
            print(f"  -> Processed & Validated {count} external portal jobs ({status})")
            if rejected:
                print(f"  [!] Rejected {len(rejected)} records due to schema invalidity")
            
            summary_stats.append({
                "source_id": scraper.source_id,
                "platform": scraper.source_name,
                "status": status,
                "jobs_count": count
            })
            
            for job in valid_jobs:
                d = job.model_dump()
                all_records.append({
                    "Platform / Portal": d.get("source_name") or scraper.source_name,
                    "Source ID": d.get("source_id") or scraper.source_id,
                    "Job Title": d.get("title") or "",
                    "Company": d.get("company") or "",
                    "Location": d.get("location") or "Pan India",
                    "Is Remote": "Yes" if d.get("is_remote") else "No",
                    "Employment Type": d.get("employment_type") or "Full Time",
                    "Salary (Scale / Raw)": d.get("salary_raw") or "Disclosed upon interview",
                    "Salary Min": d.get("salary_min") or "",
                    "Salary Max": d.get("salary_max") or "",
                    "Salary Currency": d.get("salary_currency") or "INR",
                    "Apply URL": d.get("apply_url") or "",
                    "Scraped Timestamp": d.get("scraped_at") or datetime.now().isoformat()
                })
                
        except Exception as e:
            print(f"  [ERROR] Failed to run {scraper_name}: {e}")
            summary_stats.append({
                "source_id": getattr(scraper, "source_id", "ERR"),
                "platform": getattr(scraper, "source_name", scraper_name),
                "status": "FAILED",
                "jobs_count": 0
            })
            if scraper:
                try:
                    await scraper.close()
                except Exception:
                    pass

    # Ensure output directory exists
    output_dir = r"C:\Users\ASUS\.gemini\antigravity-ide\scratch\leadgen\module_1_scraping_engine\data\output"
    os.makedirs(output_dir, exist_ok=True)
    
    csv_file = os.path.join(output_dir, "archetype_e_test_scraped_data.csv")
    xlsx_file = os.path.join(output_dir, "archetype_e_test_scraped_data.xlsx")
    
    # Save CSV
    with open(csv_file, mode="w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(all_records)
    print(f"\n[+] Saved {len(all_records)} jobs to CSV: {csv_file}")
    
    # Save Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Archetype_E_Jobs"
    ws.append(fieldnames)
    
    # Formatting header row
    header_fill = openpyxl.styles.PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = openpyxl.styles.Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        
    for rec in all_records:
        ws.append([rec.get(h, "") for h in fieldnames])
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 45)
        
    wb.save(xlsx_file)
    print(f"[+] Saved {len(all_records)} jobs to Excel: {xlsx_file}")
    
    # Print Master Summary
    print("\n" + "="*70)
    print("ARCHETYPE E TEST SUMMARY SCORECARD")
    print("="*70)
    print(f"{'Source ID':<12} | {'Portal Name':<32} | {'Status':<10} | {'Jobs Scraped'}")
    print("-" * 70)
    total_jobs = 0
    for s in summary_stats:
        total_jobs += s['jobs_count']
        print(f"{s['source_id']:<12} | {s['platform'][:32]:<32} | {s['status']:<10} | {s['jobs_count']}")
    print("-" * 70)
    print(f"TOTAL: {len(summary_stats)} Stealth Portals | {total_jobs} Valid Listings Ingested")
    print("="*70)

if __name__ == "__main__":
    asyncio.run(run_all_archetype_e())
